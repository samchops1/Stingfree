"""Falcon DLP — file / folder PII discovery scanner.

Extracts text from documents and runs the SAME Presidio engine + custom RIA
recognizers the live agent uses, then reports what PII is where. This is the
DETECTION/DISCOVERY side of Presidio pointed at documents-at-rest — useful for
auditing your document stores (which files hold client PII, and what kind).

It does NOT block, move, or modify files — enforcing on file copy/USB/email
needs an OS filter driver, which this tool deliberately doesn't build. Scan
reports in; nothing leaves the machine.

Supported: .txt .csv .tsv .md .log .json (plain text), .pdf, .docx, .xlsx.
Other types are skipped and noted.
"""

import json
import os

import detector

# Text-like extensions read directly.
TEXT_EXTS = {".txt", ".csv", ".tsv", ".md", ".log", ".json"}
# Rich types via optional libraries (imported lazily so a missing lib only
# disables that one type, and PyInstaller only bundles what's installed).
RICH_EXTS = {".pdf", ".docx", ".xlsx"}

SUPPORTED_EXTS = TEXT_EXTS | RICH_EXTS

# spaCy caps single-doc length; scan large units in chunks under that.
MAX_CHARS = 80_000

_SEVERITY = {"allow": 0, "monitor": 1, "redact": 2, "block": 3}


# --- text extraction: returns a list of (unit_label, text) -------------------

def _extract_text_file(path):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return [("file", f.read())]


def _extract_pdf(path):
    from pypdf import PdfReader  # lazy
    reader = PdfReader(path)
    units = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            units.append((f"p{i}", page.extract_text() or ""))
        except Exception:
            units.append((f"p{i}", ""))
    return units


def _extract_docx(path):
    import docx  # python-docx, lazy
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return [("document", "\n".join(parts))]


def _extract_xlsx(path):
    from openpyxl import load_workbook  # lazy
    wb = load_workbook(path, read_only=True, data_only=True)
    units = []
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows(values_only=True):
            cells.extend(str(v) for v in row if v is not None)
        units.append((f"sheet:{ws.title}", " ".join(cells)))
    wb.close()
    return units


_EXTRACTORS = {".pdf": _extract_pdf, ".docx": _extract_docx, ".xlsx": _extract_xlsx}


def extract_units(path):
    """Return [(unit_label, text), ...] for a file, or raise on unsupported."""
    ext = os.path.splitext(path)[1].lower()
    if ext in TEXT_EXTS:
        return _extract_text_file(path)
    if ext in _EXTRACTORS:
        return _EXTRACTORS[ext](path)
    raise ValueError(f"unsupported type: {ext or '(none)'}")


def _chunks(text):
    for i in range(0, len(text), MAX_CHARS):
        yield text[i:i + MAX_CHARS]


# --- scanning ----------------------------------------------------------------

def scan_file(path):
    """Scan one file. Returns a report dict (never raises for scan errors —
    they're captured in the 'error' field)."""
    report = {
        "path": path,
        "size": _safe_size(path),
        "status": "scanned",
        "action": "clean",
        "by_type": {},      # {entity_type: {"count": n, "max_score": s}}
        "hit_units": [],    # units (pages/sheets) that had findings
        "error": None,
    }
    try:
        units = extract_units(path)
    except ValueError as exc:
        report["status"] = "skipped"
        report["error"] = str(exc)
        return report
    except Exception as exc:
        report["status"] = "error"
        report["error"] = f"extract failed: {exc}"
        return report

    worst = "clean"
    for label, text in units:
        if not text or not text.strip():
            continue
        unit_had_hit = False
        for chunk in _chunks(text):
            try:
                res = detector.scan(chunk)
            except Exception as exc:
                report["status"] = "error"
                report["error"] = f"scan failed on {label}: {exc}"
                continue
            for f in res["findings"]:
                unit_had_hit = True
                bt = report["by_type"].setdefault(f["type"], {"count": 0, "max_score": 0.0})
                bt["count"] += 1
                bt["max_score"] = max(bt["max_score"], f["score"])
            if _SEVERITY.get(res["action"], 0) > _SEVERITY.get(worst if worst != "clean" else "allow", 0):
                worst = res["action"]
        if unit_had_hit:
            report["hit_units"].append(label)

    report["action"] = worst
    return report


def scan_path(root, progress=None):
    """Scan a file or walk a folder. Returns (reports, summary)."""
    detector.warm_up()
    targets = _collect(root)
    reports = []
    for i, path in enumerate(targets, start=1):
        if progress:
            progress(i, len(targets), path)
        reports.append(scan_file(path))
    return reports, summarize(reports)


def _collect(root):
    if os.path.isfile(root):
        return [root]
    found = []
    for dirpath, _dirs, files in os.walk(root):
        for name in files:
            if os.path.splitext(name)[1].lower() in SUPPORTED_EXTS:
                found.append(os.path.join(dirpath, name))
    return sorted(found)


def _safe_size(path):
    try:
        return os.path.getsize(path)
    except OSError:
        return 0


def summarize(reports):
    scanned = [r for r in reports if r["status"] == "scanned"]
    with_pii = [r for r in scanned if r["by_type"]]
    totals = {}
    for r in with_pii:
        for t, v in r["by_type"].items():
            totals[t] = totals.get(t, 0) + v["count"]
    return {
        "files_total": len(reports),
        "files_scanned": len(scanned),
        "files_skipped": sum(1 for r in reports if r["status"] == "skipped"),
        "files_error": sum(1 for r in reports if r["status"] == "error"),
        "files_with_pii": len(with_pii),
        "entity_totals": dict(sorted(totals.items(), key=lambda kv: -kv[1])),
    }
